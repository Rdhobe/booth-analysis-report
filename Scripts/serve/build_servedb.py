"""Serving store: state CSVs + report workbook -> ONE indexed SQLite DB.

Reads (all local, offline):
- data/processed/state_booth.csv  (booth x election grain + M/F/O)
- data/processed/state_fact.csv   (booth x candidate votes, ~3.4M rows)
- data/processed/state_coverage.csv (SEG-CHECK / DEMO-CHECK / gaps)
- data/assembly-segment-level-voting-data-for-all-constituencies-2024.csv
  (ECI official 2024 AC-level candidate totals - full 403-AC fallback
  where 2024 booth conversion is thin)
- output/up_booth_report.xlsx     (All_ACs display rows = single source
  of truth for shares/zones/turnout/demo columns; Zone_Counts; Methods)
- configs/ac_format.json          (AC names, districts)

Writes: output/servedb.sqlite (WAL, indexed, FTS search, pre-aggregated
ac_summary). The FastAPI server reads this file read-only; the 272 MB
fact CSV and 59 MB xlsx are never touched at serve time.

Usage (heavy: ~5 min - run in an external terminal):
  python Scripts/build_servedb.py [--db output/servedb.sqlite]
      [--districts Agra Saharanpur]   (pilot slice for testing)
"""
from __future__ import annotations

import argparse
import csv
import json
import logging
import re
import sqlite3
import sys
import time
from pathlib import Path

logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
log = logging.getLogger("servedb")

ROOT = Path(__file__).parent.parent
DATA = ROOT / "data"
PROC = DATA / "processed"

ELECTIONS = ["VS2017", "LS2019", "VS2022", "LS2024"]

SCHEMA = """
CREATE TABLE booth_row(
  ac INTEGER, ps TEXT, district TEXT, name TEXT, address TEXT,
  s2017 TEXT, s2019 TEXT, s2022 TEXT, s2024 TEXT,
  zone TEXT, avg_diff REAL, latest_diff REAL, party_b TEXT,
  a2017 REAL, b2017 REAL, d2017 REAL,
  a2019 REAL, b2019 REAL, d2019 REAL,
  a2022 REAL, b2022 REAL, d2022 REAL,
  a2024 REAL, b2024 REAL, d2024 REAL,
  e2017 INTEGER, m2017 INTEGER, f2017 INTEGER, o2017 INTEGER,
  t2017 INTEGER, to2017 REAL,
  e2019 INTEGER, m2019 INTEGER, f2019 INTEGER, o2019 INTEGER,
  t2019 INTEGER, to2019 REAL,
  e2022 INTEGER, m2022 INTEGER, f2022 INTEGER, o2022 INTEGER,
  t2022 INTEGER, to2022 REAL,
  e2024 INTEGER, m2024 INTEGER, f2024 INTEGER, o2024 INTEGER,
  t2024 INTEGER, to2024 REAL, ac_electors_2024 INTEGER,
  {demo_cols}
  PRIMARY KEY (ac, ps));
CREATE TABLE booth_election(
  booth_uid TEXT PRIMARY KEY, election TEXT, ac INTEGER, ps TEXT,
  ps_name TEXT, electors INTEGER, male INTEGER, female INTEGER,
  other INTEGER, epic INTEGER, tendered INTEGER, polled INTEGER,
  valid INTEGER, nota INTEGER, src TEXT);
CREATE TABLE vote(
  booth_uid TEXT, candidate TEXT, party TEXT, votes INTEGER,
  PRIMARY KEY (booth_uid, candidate));
CREATE TABLE ac_summary(
  ac INTEGER, election TEXT, ac_name TEXT, district TEXT,
  booths INTEGER, electors INTEGER, polled INTEGER, valid INTEGER,
  nota INTEGER, turnout_pct REAL,
  PRIMARY KEY (ac, election));
CREATE TABLE ac_party(
  ac INTEGER, election TEXT, party TEXT, votes INTEGER,
  PRIMARY KEY (ac, election, party));
CREATE TABLE ac_zones(
  ac INTEGER, zone TEXT, n INTEGER,
  PRIMARY KEY (ac, zone));
CREATE TABLE coverage(
  id INTEGER PRIMARY KEY AUTOINCREMENT, election TEXT, ac TEXT,
  ac_name TEXT, file TEXT, booths TEXT, valid_ok TEXT, has_total TEXT,
  note TEXT);
CREATE TABLE seg2024(
  ac INTEGER, candidate TEXT, party TEXT, votes INTEGER,
  PRIMARY KEY (ac, candidate));
CREATE TABLE seg2024_ac(
  ac INTEGER PRIMARY KEY, electors INTEGER, nota INTEGER,
  valid_sum INTEGER);
CREATE TABLE meta(key TEXT PRIMARY KEY, value TEXT);
CREATE VIRTUAL TABLE booth_fts USING fts5(
  ac, ps, name, district,
  content='booth_row', content_rowid='rowid');
"""


def _int(x):
    if x is None or (isinstance(x, str) and not x.strip()):
        return None
    try:
        return int(float(str(x).replace(",", "")))
    except (TypeError, ValueError):
        return None


def _real(x):
    if x is None or (isinstance(x, str) and not x.strip()):
        return None
    try:
        return float(str(x).replace(",", ""))
    except (TypeError, ValueError):
        return None


def load_ac_info():
    d = json.loads((ROOT / "configs" / "ac_format.json").read_text(
        encoding="utf-8"))
    return {int(a["ac_no"]): (a.get("ac_name", ""),
                               a.get("district", "UNKNOWN"))
            for a in d["assembly_constituencies"]}


def load_seg2024(want_ac=None):
    """ECI assembly-segment file (official 2024 AC-level results).

    Title row + header + one row per (AC, candidate); UP rows only.
    Returns (cand_rows, ac_stats): cand_rows = [(ac, candidate, party,
    votes)], ac_stats = {ac: {electors, nota, valid_sum}}."""
    p = (DATA / "assembly-segment-level-voting-data-for-all-"
         "constituencies-2024.csv")
    with open(p, encoding="utf-8-sig") as f:
        lines = [ln for ln in f if ln.strip()]
    if lines and not lines[0].startswith("State/UT"):
        lines = lines[1:]  # ECI title row, not data
    cand_rows: list = []
    acs: dict = {}
    for r in csv.DictReader(lines):
        if (r.get("State/UT Name") or "").strip() != "Uttar Pradesh":
            continue
        try:
            ac = int(r.get("AC NO"))
        except (TypeError, ValueError):
            continue
        if want_ac is not None and ac not in want_ac:
            continue
        st = acs.setdefault(ac, {"electors": _int(
            r.get("TOTAL ELECTORS IN AC")), "nota": _int(
            r.get("NOTA VOTES EVM IN AC")), "valid_sum": 0})
        v = _int(r.get("VOTES SECURED EVM"))
        st["valid_sum"] += v or 0
        name = (r.get("CANDIDATE NAME") or "").strip()
        party = (r.get("PARTY") or "").strip().upper()
        if name and party and v is not None:
            cand_rows.append((ac, name, party, v))
    return cand_rows, acs


def read_allacs(xlsx_path, want_ac=None):
    """All_ACs rows positionally (robust to header renames).

    Layout: A AC, B District, C name, D address, E-H serials, I zone,
    J avg, K latest, L blank, M party_b, then per election 5 cols
    [A%, B%, blank, diff, blank] x4, then 24 turnout cols
    [E,M,F,O,T,TO%]x4, then 2024 AC electors, then 36 demo cols.
    Returns (rows, demo_cols). Data starts at row 4 (1-indexed)."""
    import openpyxl
    wb = openpyxl.load_workbook(str(xlsx_path), read_only=True,
                                data_only=True)
    ws = wb["All_ACs"]
    demo_cols: list = []
    out: list = []
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i % 50000 == 0:
            log.info("scanning All_ACs row %d ...", i)
        if i == 0:
            continue  # title
        if i == 1:
            hdr = [str(v or "") for v in r]
            tail = hdr[58:]
            demo_cols = [re.sub(r"\s*\(modeled\)\s*$", "", c).strip()
                         for c in tail]
            assert len(r) == 94, f"All_ACs width changed: {len(r)}"
            continue
        if i == 2:
            continue  # Hindi header
        if r[0] in (None, ""):
            continue
        try:
            ac = int(r[0])
        except (TypeError, ValueError):
            continue
        if want_ac is not None and ac not in want_ac:
            continue
        vals = list(r) + [None] * (94 - len(list(r)))
        shares = []
        for b in range(4):
            o = 13 + b * 5
            shares += [_real(vals[o]), _real(vals[o + 1]),
                       _real(vals[o + 3])]
        to = []
        for b in range(4):
            o = 33 + b * 6
            e, m, f, oo, t, top = vals[o:o + 6]
            to += [_int(e), _int(m), _int(f), _int(oo), _int(t),
                   _real(top)]
        rec = {
            "ac": ac, "ps": str(vals[2]).split(" - ", 1)[0]
            if vals[2] else "",
            "district": vals[1] or "", "name": vals[2] or "",
            "address": vals[3] or "",
            "serials": [vals[4], vals[5], vals[6], vals[7]],
            "zone": vals[8] or "", "avg": _real(vals[9]),
            "latest": _real(vals[10]), "party_b": vals[12] or "",
            "shares": shares, "to": to,
            "ac_electors_2024": _int(vals[57]),
            "demo": [_real(v) if not isinstance(v, str)
                     or re.match(r"^-?[\d.,]+$", v.strip()) else None
                     for v in vals[58:94]],
        }
        # keep raw demo strings that aren't plain numbers as NULL
        rec["demo"] = [v if isinstance(v, (int, float)) else None
                       for v in rec["demo"]]
        out.append(rec)
    wb.close()
    return out, demo_cols


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Build serving SQLite DB")
    ap.add_argument("--db", default="output/servedb.sqlite")
    ap.add_argument("--districts", nargs="*", default=[],
                    help="pilot slice: only these districts' ACs are "
                         "loaded (else full 403-AC build). NOTE: source "
                         "files are still fully parsed - this shrinks "
                         "the DB, not the runtime.")
    ap.add_argument("--report",
                    default="output/up_booth_report.xlsx")
    args = ap.parse_args(argv)
    t_all = time.time()

    ac_info = load_ac_info()
    want_ac = None
    if args.districts:
        want = {d.lower().strip() for d in args.districts}
        all_d = {d.lower() for _a, d in ac_info.values()}
        unknown = want - all_d
        if unknown:
            log.error("unknown district(s): %s", ", ".join(sorted(unknown)))
            log.error("known: %s", ", ".join(sorted(all_d)))
            return 2
        want_ac = {ac for ac, (_n, d) in ac_info.items()
                   if d.lower() in want}
        if not want_ac:
            log.error("district filter matched 0 ACs - aborting")
            return 2
        log.info("pilot districts=%s -> %d of %d ACs", args.districts,
                 len(want_ac), len(ac_info))

    dbp = ROOT / args.db
    if dbp.exists():
        dbp.unlink()
    dbp.parent.mkdir(parents=True, exist_ok=True)

    log.info("reading All_ACs (single source of truth) ...")
    rows, demo_cols = read_allacs(ROOT / args.report, want_ac)
    log.info("booth_rows=%d demo_cols=%d", len(rows), len(demo_cols))

    con = sqlite3.connect(str(dbp))
    con.execute("PRAGMA journal_mode=WAL;")
    con.execute("PRAGMA synchronous=OFF;")
    demo_ddl = ",\n  ".join(f'"{c}" REAL' for c in demo_cols)
    con.executescript(SCHEMA.format(demo_cols=demo_ddl + ","
                                    if demo_ddl else ""))

    # ---- booth_row + FTS
    ins_cols = ("ac, ps, district, name, address, s2017, s2019, s2022, "
                "s2024, zone, avg_diff, latest_diff, party_b, "
                "a2017, b2017, d2017, a2019, b2019, d2019, "
                "a2022, b2022, d2022, a2024, b2024, d2024, "
                "e2017, m2017, f2017, o2017, t2017, to2017, "
                "e2019, m2019, f2019, o2019, t2019, to2019, "
                "e2022, m2022, f2022, o2022, t2022, to2022, "
                "e2024, m2024, f2024, o2024, t2024, to2024, "
                "ac_electors_2024" + (", " + ", ".join(
                    f'"{c}"' for c in demo_cols) if demo_cols else ""))
    n_ph = ins_cols.count(",") + 1
    cur = con.cursor()
    batch: list = []
    n_ins = 0
    for r in rows:
        vals = [r["ac"], r["ps"], r["district"], r["name"],
                r["address"]] + [str(s) if s not in (None, "") else None
                                          for s in r["serials"]] + [
            r["zone"], r["avg"], r["latest"], r["party_b"]] + r[
            "shares"] + r["to"] + [r["ac_electors_2024"]] + r["demo"]
        assert len(vals) == n_ph, (len(vals), n_ph)
        batch.append(tuple(vals))
        if len(batch) >= 2000:
            cur.executemany(f"INSERT INTO booth_row ({ins_cols}) VALUES "
                            f"({','.join('?' * n_ph)})", batch)
            n_ins += len(batch)
            batch = []
            if n_ins % 20000 < 2000:
                log.info("booth_row ... %d / %d", n_ins, len(rows))
    if batch:
        cur.executemany(f"INSERT INTO booth_row ({ins_cols}) VALUES "
                        f"({','.join('?' * n_ph)})", batch)
    con.commit()
    log.info("booth_row done (%.0fs)", time.time() - t_all)
    cur.execute("INSERT INTO booth_fts(ac, ps, name, district) "
                "SELECT ac, ps, name, district FROM booth_row")
    con.commit()

    # ---- booth_election (state_booth.csv)
    n_be = 0
    batch = []
    with open(PROC / "state_booth.csv", encoding="utf-8") as f:
        for b in csv.DictReader(f):
            try:
                ac = int(b["ac"])
            except (TypeError, ValueError):
                continue
            if want_ac is not None and ac not in want_ac:
                continue
            batch.append((b["booth_uid"], b["election"], ac, b["ps"],
                          b.get("ps_name", ""), _int(b.get("electors")),
                          _int(b.get("male")), _int(b.get("female")),
                          _int(b.get("other")), _int(b.get("epic")),
                          _int(b.get("tendered")), _int(b.get("polled")),
                          _int(b.get("valid")), _int(b.get("nota")),
                          b.get("src", "")))
            if len(batch) >= 5000:
                cur.executemany(
                    "INSERT OR IGNORE INTO booth_election VALUES "
                    "(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", batch)
                n_be += len(batch)
                batch = []
                if n_be % 100000 < 5000:
                    log.info("booth_election ... %d", n_be)
    if batch:
        cur.executemany("INSERT OR IGNORE INTO booth_election VALUES "
                        "(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", batch)
        n_be += len(batch)
    con.commit()
    log.info("booth_election=%d (%.0fs)", n_be, time.time() - t_all)

    # ---- vote (state_fact.csv; ~3.4M rows)
    n_v = 0
    batch = []
    with open(PROC / "state_fact.csv", encoding="utf-8") as f:
        for v in csv.DictReader(f):
            try:
                ac = int(v["ac"])
            except (TypeError, ValueError):
                continue
            if want_ac is not None and ac not in want_ac:
                continue
            try:
                vv = int(v["votes"])
            except (TypeError, ValueError):
                continue
            batch.append((v["booth_uid"], v["candidate"],
                          v["party"], vv))
            if len(batch) >= 20000:
                cur.executemany("INSERT OR IGNORE INTO vote VALUES "
                                "(?,?,?,?)", batch)
                n_v += len(batch)
                batch = []
                if n_v % 500000 < 20000:
                    log.info("vote ... %d", n_v)
    if batch:
        cur.executemany("INSERT OR IGNORE INTO vote VALUES (?,?,?,?)",
                        batch)
        n_v += len(batch)
    con.commit()
    log.info("vote=%d (%.0fs)", n_v, time.time() - t_all)

    # ---- seg2024 (official 2024 AC totals; fallback where booth
    # conversion is thin)
    seg_rows, seg_acs = load_seg2024(want_ac)
    cur.executemany("INSERT OR IGNORE INTO seg2024 VALUES (?,?,?,?)",
                    seg_rows)
    cur.executemany("INSERT OR REPLACE INTO seg2024_ac VALUES (?,?,?,?)",
                    [(ac, s["electors"], s["nota"], s["valid_sum"])
                     for ac, s in seg_acs.items()])
    con.commit()
    log.info("seg2024=%d acs (%d) (%.0fs)", len(seg_rows), len(seg_acs),
             time.time() - t_all)

    # ---- coverage
    with open(PROC / "state_coverage.csv", encoding="utf-8") as f:
        for c in csv.DictReader(f):
            cur.execute("INSERT INTO coverage (election, ac, ac_name, "
                        "file, booths, valid_ok, has_total, note) VALUES "
                        "(?,?,?,?,?,?,?,?)",
                        (c.get("election", ""), c.get("ac", ""),
                         c.get("ac_name", ""), c.get("file", ""),
                         c.get("booths", ""), c.get("valid_ok", ""),
                         c.get("has_total", ""), c.get("note", "")))
    con.commit()

    # ---- pre-aggregations (SQL, fast)
    log.info("aggregating ...")
    cur.execute("INSERT INTO ac_summary "
                "SELECT ac, election, '', '', COUNT(*), "
                "SUM(electors), SUM(polled), SUM(valid), SUM(nota), "
                "CASE WHEN SUM(electors) > 0 THEN ROUND("
                "100.0 * SUM(polled) / SUM(electors), 2) END "
                "FROM booth_election GROUP BY ac, election")
    for ac, (nm, dist) in ac_info.items():
        cur.execute("UPDATE ac_summary SET ac_name=?, district=? "
                    "WHERE ac=?", (nm, dist, ac))
    cur.execute("INSERT INTO ac_party "
                "SELECT b.ac, b.election, v.party, SUM(v.votes) "
                "FROM vote v JOIN booth_election b "
                "ON v.booth_uid = b.booth_uid "
                "GROUP BY b.ac, b.election, v.party")
    cur.execute("INSERT INTO ac_zones SELECT ac, zone, COUNT(*) "
                "FROM booth_row GROUP BY ac, zone")
    # methods text (honesty block, verbatim from workbook)
    import openpyxl
    wb = openpyxl.load_workbook(str(ROOT / args.report), read_only=True,
                                data_only=True)
    methods = wb["Methods_Zones"]["A1"].value or ""
    zones = [[c.value for c in row]
             for row in wb["Zone_Counts"].iter_rows()]
    wb.close()
    cur.execute("INSERT INTO meta VALUES (?,?)",
                ("methods", methods))
    cur.execute("INSERT INTO meta VALUES (?,?)",
                ("zone_counts", json.dumps(zones)))
    cur.execute("INSERT INTO meta VALUES (?,?)",
                ("demo_cols", json.dumps(demo_cols)))
    cur.execute("INSERT INTO meta VALUES (?,?)",
                ("built_at", time.strftime("%Y-%m-%d %H:%M:%S")))
    cur.execute("INSERT INTO meta VALUES (?,?)",
                 ("elections", json.dumps(ELECTIONS)))
    cur.execute("INSERT INTO meta VALUES (?,?)",
                 ("scope", ",".join(args.districts) if args.districts
                  else "statewide (all %d ACs)" % len(ac_info)))
    cur.execute("INSERT INTO meta VALUES (?,?)",
                 ("slice", "1" if args.districts else "0"))
    con.commit()

    # ---- indexes (after load = fast bulk path)
    cur.executescript("""
CREATE INDEX idx_be_ac ON booth_election(ac, election);
CREATE INDEX idx_be_uid ON booth_election(booth_uid);
CREATE INDEX idx_vote_uid ON vote(booth_uid);
CREATE INDEX idx_vote_party ON vote(party);
CREATE INDEX idx_row_ac ON booth_row(ac);
CREATE INDEX idx_row_zone ON booth_row(zone);
CREATE INDEX idx_row_district ON booth_row(district);
CREATE INDEX idx_row_zone_a2017 ON booth_row(zone, a2017);
CREATE INDEX idx_row_zone_a2019 ON booth_row(zone, a2019);
CREATE INDEX idx_row_zone_a2022 ON booth_row(zone, a2022);
CREATE INDEX idx_row_zone_a2024 ON booth_row(zone, a2024);
ANALYZE;
""")
    con.execute("PRAGMA synchronous=FULL;")
    con.commit()
    n = {t: cur.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0]
         for t in ("booth_row", "booth_election", "vote", "ac_summary",
                   "ac_party", "coverage", "seg2024", "seg2024_ac")}
    con.close()
    log.info("BUILD COMPLETE counts=%s total=%.0fs -> %s",
             n, time.time() - t_all, dbp)
    return 0


if __name__ == "__main__":
    sys.exit(main())
