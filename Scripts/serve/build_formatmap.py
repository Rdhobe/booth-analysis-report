"""Map every Excel file's structure -> configs/data_file_format.json.

Per-file entries (user call): sheet list with dims, header/data row
positions, and per-column {index, name, dtype} where dtype is text or
numbers. No sample values. PDFs are left exactly as mapped in
configs/data_files.json (no structure attempted).

Layouts:
- form20-xlsx: our converter output (Form20 + Validation sheets).
  Title rows 1-6, merged header row 7, names row 8, data from row 9:
  col0 serial (numbers), col1 PS no/name (text), candidate vote cols
  (numbers), last 5 trailers valid|rejected|NOTA|total|tendered (numbers).
- eci-triplet: 2017/2019/2022 ECI sheets. Booth block cols 0-9
  (AC, PS-No, PS-Name, Electors, Male, Female, Other, Turnout-Total,
  EPIC, Tendered), then per-candidate [S.No|Party|Votes] triplets
  (name text, party text, votes numbers), then Total/NOTA trailers.

Usage: python Scripts/build_formatmap.py
Output: configs/data_file_format.json
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).parent.parent
MAP_IN = ROOT / "configs" / "data_files.json"
OUT = ROOT / "configs" / "data_file_format.json"

BOOTH_COLS = [("AC No.", "numbers"), ("PS No.", "numbers"),
              ("PS Name", "text"), ("Electors", "numbers"),
              ("Turnout Male", "numbers"), ("Turnout Female", "numbers"),
              ("Turnout Other", "numbers"), ("Turnout Total", "numbers"),
              ("EPIC identified", "numbers"), ("Tendered", "numbers")]


def clean(s) -> str:
    return re.sub(r"\s+", " ", str(s or "").strip())


def map_form20_xlsx(path: Path):
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    sheets = []
    for ws in wb.worksheets:
        if ws.title == "Validation":
            sheets.append({"name": "Validation", "nrows": ws.max_row,
                           "ncols": ws.max_column, "header_rows": [],
                           "data_start_row": None, "columns": [],
                           "note": "converter validation log (not results)"})
            continue
        # candidate names row 8, trailers = last 5 cols
        ncols = ws.max_column or 0
        names = [clean(ws.cell(8, c).value) for c in range(3, ncols - 4)]
        cols = [{"index": 0, "name": "Serial No.", "dtype": "numbers"},
                {"index": 1, "name": "Polling Station No./Name", "dtype": "text"}]
        for i, nm in enumerate(names):
            cols.append({"index": 2 + i, "name": nm or f"candidate_{i + 1}",
                         "dtype": "numbers"})
        for j, nm in enumerate(["Total Valid Votes", "Rejected Votes",
                                "NOTA", "Total", "Tendered Votes"]):
            cols.append({"index": 2 + len(names) + j, "name": nm,
                         "dtype": "numbers"})
        sheets.append({"name": ws.title, "nrows": ws.max_row, "ncols": ncols,
                       "header_rows": [7, 8], "data_start_row": 9,
                       "columns": cols, "note": ""})
    wb.close()
    return sheets, "form20-xlsx"


def map_eci_xls(path: Path):
    import xlrd
    bk = xlrd.open_workbook(str(path))
    sheets = []
    for si in range(bk.nsheets):
        sh = bk.sheet_by_index(si)
        # vote-header row = most 'Votes Secured' hits
        best, hr = None, -1
        for r in range(min(12, sh.nrows)):
            cols = [c for c in range(sh.ncols)
                    if "Votes Secured" in str(sh.cell(r, c).value)]
            if cols and (best is None or len(cols) > len(best[1])):
                best = (r, cols)
        if best is None:
            sheets.append({"name": sh.name, "nrows": sh.nrows,
                           "ncols": sh.ncols, "header_rows": [],
                           "data_start_row": None, "columns": [],
                           "note": "no candidate triplets (check/round sheet?)"})
            continue
        hr, vcols = best
        # trailing Total/NOTA zone (not candidates)
        lab = lambda cc: " ".join(  # noqa: E731
            str(sh.cell(rr, cc).value or "") for rr in {hr - 1, hr} if rr >= 0)
        tz = next((cc for cc in range(10, sh.ncols)
                   if re.search(r"total|tendered|reject", lab(cc), re.I)), None)
        if tz is not None:
            vcols = [v for v in vcols if v < tz]
        cols = [{"index": i, "name": n, "dtype": d}
                for i, (n, d) in enumerate(BOOTH_COLS)]
        # candidate names live one row above (nearest wins)
        for v in vcols:
            nm = ""
            for nr in (hr - 1, hr - 2):
                if nr < 0:
                    continue
                for cc in range(max(0, v - 2), min(v + 1, sh.ncols)):
                    t = clean(sh.cell(nr, cc).value)
                    if t and not re.match(
                            r"^(S\.?No\.?|Party Affiliation|Votes Secured|No\.|Name|Total)$",
                            t, re.I):
                        nm = re.sub(r"^\d+\s*[-–.]\s*", "", t)
                        break
                if nm:
                    break
            base = len(cols)
            cols += [{"index": base, "name": f"{nm or 'candidate'} :: S.No",
                      "dtype": "text"},
                     {"index": base + 1,
                      "name": f"{nm or 'candidate'} :: Party Affiliation",
                      "dtype": "text"},
                     {"index": base + 2,
                      "name": f"{nm or 'candidate'} :: Votes Secured",
                      "dtype": "numbers"}]
        # explicit trailing Total/NOTA cols
        for cc in range(10, sh.ncols):
            lab = " ".join(str(sh.cell(rr, cc).value or "")
                           for rr in ({hr - 1, hr} - {-1})).strip()
            if re.match(r"^\s*NOTA\s*$", lab, re.I):
                cols.append({"index": cc, "name": "NOTA", "dtype": "numbers"})
            elif re.search(r"total", lab, re.I) and \
                    not any(v == cc for v in vcols) and \
                    not any(c["index"] == cc for c in cols):
                cols.append({"index": cc, "name": lab[:40] or "Total",
                             "dtype": "numbers"})
        # data start: first row below header with numeric PS col
        data_start = None
        for r in range(hr + 1, sh.nrows):
            c1 = str(sh.cell(r, 1).value).strip() if sh.ncols > 1 else ""
            if re.match(r"^\d+(\.0)?$", c1):
                data_start = r + 1  # 1-indexed
                break
        sheets.append({"name": sh.name, "nrows": sh.nrows, "ncols": sh.ncols,
                       "header_rows": [hr, hr + 1], "data_start_row": data_start,
                       "columns": sorted(cols, key=lambda c: c["index"]),
                       "note": ""})
    # layout label
    kind = "eci-triplet"
    return sheets, kind


def main() -> int:
    doc = json.loads(MAP_IN.read_text(encoding="utf-8"))
    by_src = {}
    for r in doc["files"]:
        by_src[r.get("source", "")] = r
    out_files, counts = [], {}
    # every excel file on disk (xls + converted/review xlsx sidecars)
    targets = sorted((ROOT / "data").rglob("*.xls")) + \
        sorted((ROOT / "data").rglob("*.xlsx"))
    # skip non-Form20 reference workbooks at data/ top level
    targets = [p for p in targets if p.parent != ROOT / "data"]
    print(f"excel files on disk: {len(targets)}", flush=True)
    for i, path in enumerate(targets, start=1):
        rel = str(path.relative_to(ROOT)).replace("\\", "/")
        base = by_src.get(rel, {})
        try:
            if path.suffix == ".xlsx":
                sheets, layout = map_form20_xlsx(path)
            else:
                sheets, layout = map_eci_xls(path)
            status = "mapped"
        except Exception as ex:  # noqa: BLE001
            sheets, layout, status = [], "unreadable", f"error: {ex}"
        else:
            status = "mapped"
        entry = {"file": rel,
                 "year": base.get("year", ""),
                 "district": base.get("district", ""),
                 "ac_no": base.get("ac_no"),
                 "layout": layout,
                 "status": status,
                 "sheets": sheets}
        if status != "mapped":
            entry["note"] = status
        out_files.append(entry)
        counts[layout if status == "mapped" else "unreadable"] = \
            counts.get(layout if status == "mapped" else "unreadable", 0) + 1
        if i % 300 == 0:
            print(f"  {i}/{len(targets)}", flush=True)
    OUT.write_text(json.dumps({"meta": {
        "generated_by": "Scripts/build_formatmap.py",
        "from": "configs/data_files.json + on-disk Excel (PDFs untouched)",
        "column": {"index": "0-based position in sheet",
                   "name": "header text as-is (serial prefix stripped)",
                   "dtype": "text = names/labels/parties; numbers = counts/votes"}},
        "files": out_files}, ensure_ascii=False, indent=1), encoding="utf-8")
    print("wrote", OUT, counts)
    return 0


if __name__ == "__main__":
    sys.exit(main())
