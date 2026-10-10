"""Independent check for PDFs that carry an embedded text layer (e.g. the Hindi sheets).

    python eval/text_layer_check.py data/input/8Hin.pdf data/output/8Hin_Form20.xlsx

Every number found in the PDF's own text layer is located in the table grid and compared with
the value in the produced workbook. The text layer is partial (often ~25% of the cells) and
not always right, so this is a sanity check, not a full ground truth.
"""
import argparse
import bisect
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import fitz  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from form20.grid import apply_template, build_col_template, detect_lines  # noqa: E402
from form20.imaging import TARGET_WIDTH, load_page, page_count  # noqa: E402


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pdf")
    ap.add_argument("xlsx")
    a = ap.parse_args()
    ws = load_workbook(a.xlsx)["Form20"]
    hdr = next(r for r in range(1, 30) if ws.cell(row=r, column=1).value == "Serial No.")
    heads = [ws.cell(row=hdr, column=c).value for c in range(1, ws.max_column + 1)]
    c_first = 3
    n_cand = 0
    while heads[c_first - 1 + n_cand] is not None:     # candidate names run until the first empty header
        n_cand += 1
    c_last = c_first + n_cand + 4                         # + valid, rejected, NOTA, total, tendered
    c_page, c_row = heads.index("Source_Page") + 1, heads.index("Source_Row") + 1
    vals = {}
    for r in range(hdr + 1, ws.max_row + 1):
        if isinstance(ws.cell(row=r, column=1).value, int):
            vals[(ws.cell(row=r, column=c_page).value, ws.cell(row=r, column=c_row).value)] = [
                ws.cell(row=r, column=c).value for c in range(c_first, c_last + 1)
            ]
    n = page_count(a.pdf)
    pls = [detect_lines(load_page(a.pdf, p)[0]) for p in range(n)]
    tpl = build_col_template(pls)
    doc = fitz.open(a.pdf)
    id_cols = len(tpl) - 1 - (c_last - c_first + 1)
    total = agree = 0
    bad = []
    for p in range(n):
        grid = apply_template(pls[p], tpl)
        scale = TARGET_WIDTH / doc[p].rect.width
        first = 1 if p == 0 and id_cols == 1 else 0     # Hindi page 1 has the 1..19 index row
        for x0, y0, x1, y1, txt, *_ in doc[p].get_text("words"):
            if not txt.isdigit():
                continue
            cx, cy = (x0 + x1) / 2 * scale, (y0 + y1) / 2 * scale
            ri = bisect.bisect_right(grid.data_edges, cy) - 1
            ci = bisect.bisect_right(grid.col_edges, cx) - 1
            if ri < first or ri >= len(grid.data_edges) - 1 or ci < id_cols or ci >= len(grid.col_edges) - 1:
                continue
            v = vals.get((p + 1, ri - first + 1))
            if not v:
                continue
            total += 1
            if str(v[ci - id_cols]) == txt:
                agree += 1
            else:
                bad.append((p + 1, ri - first + 1, ci - id_cols, txt, v[ci - id_cols]))
    print(f"numbers in the PDF text layer that fall in the table: {total}")
    print(f"identical to the workbook: {agree} ({100 * agree / max(1, total):.2f}%)")
    for b in bad[:30]:
        print("  page %d row %d col %d: text layer %s, workbook %s" % b)
    return 0 if not bad else 1


if __name__ == "__main__":
    sys.exit(main())
