"""Stage verified OCR workbooks in the layout of data/2024/EXCEL, named <AC>.xlsx.

    python batch/stage_verified.py [--out data/output/staged_excel]

Only workbooks that the official check marks FULL (every candidate column and NOTA equal the official 2024 EVM totals)
are staged. Candidate names come from the official file: each sheet column is matched to the official candidate with
the same total (ties resolved by name similarity). Nothing is written to data/2024/EXCEL - copy the files over by hand
once you are happy with them.
"""
import argparse
import sys
from collections import OrderedDict
from difflib import SequenceMatcher
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from queue_common import QUEUE_DIR, TOOL_DIR, read_csv  # noqa: E402
from verify_official import OFFICIAL_CSV, load_official  # noqa: E402

DEFAULT_OUT = TOOL_DIR / "data" / "output" / "staged_excel"
TOTAL_METRICS = ["TOTAL_VALID", "REJECTED", "NOTA", "TOTAL", "TENDERED"]


def read_tidy(xlsx_path):
    """Booth rows from a workbook's Tidy sheet -> (labels in sheet order, [(serial, ps_no, [votes per candidate
    column], {metric: votes})], electors). Columns are told apart by position: OCR often repeats names."""
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    labels, booths, electors = [], [], None
    tidy = list(wb["Tidy"].iter_rows(values_only=True))[1:]
    wb.close()                                    # release the file (Windows keeps read-only workbooks locked)
    prev, cur = None, None
    for r in tidy:
        serial, ps_no, metric, name, votes = r[4], r[5], r[6], r[7], r[10]
        electors = electors or r[3]
        if metric == "CANDIDATE" and prev != "CANDIDATE":     # a new booth (serials restart in multi-AC PDFs)
            cur = (serial, ps_no, [], {})
            booths.append(cur)
        if metric == "CANDIDATE":
            if len(cur[2]) >= len(labels):
                labels.append(name)
            cur[2].append(votes)
        else:
            cur[3][metric] = votes
        prev = metric
    return labels, booths, electors


def assign_names(labels, sums, official_candidates):
    """Match every sheet column to the official candidate with the same total; equal totals are told apart by
    similarity of the OCR label to the official name. Returns names in sheet order (the OCR label when unmatched)."""
    free = list(official_candidates)                        # (name, party, votes)
    names = {}
    order = sorted(range(len(labels)), key=lambda i: sum(1 for c in free if c[2] == sums[i]))   # rarest totals first
    for i in order:
        same = [c for c in free if c[2] == sums[i]]
        if not same:
            names[i] = (str(labels[i]), "")
            continue
        best = max(same, key=lambda c: SequenceMatcher(None, str(labels[i]).upper(), c[0].upper()).ratio())
        free.remove(best)
        names[i] = (best[0], best[1])
    return [names[i] for i in range(len(labels))]


def build_workbook(ac, source_name, labels, booths, names, official_ac, out_path, swap_nota=False, verified=True,
                   note=""):
    """Write the legacy-layout workbook. booths: (serial, ps_no, [votes per candidate column], {metric: votes});
    names: [(name, party)] in sheet order."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter as L

    nc = len(labels)
    ncols = 2 + nc + 5
    wb = Workbook()
    ws = wb.active
    ws.title = "Form20"
    title = ["FORM 20", "FINAL RESULT SHEET", "ELECTION TO THE HOUSE OF PARLIAMENT CONSTITUENCY",
             f"Total No. of Electors in Assembly Constituency/segment ....{official_ac.get('electors', '')}",
             f"Name of Assembly/segment ...{ac}-{official_ac['name']} Assembly Election",
             f"Source file: {source_name} (OCR, " + ("verified against the official EVM totals)" if verified
                                                   else "NOT VERIFIED against the official EVM totals)")]
    bold = Font(bold=True)
    for i, t in enumerate(title, 1):
        ws.cell(i, 1, t).font = bold
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=ncols)
        ws.cell(i, 1).alignment = Alignment(horizontal="center")
    fill = PatternFill("solid", fgColor="D9E1F2")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.cell(7, 1, "Serial No. Of Polling Station")
    ws.merge_cells(start_row=7, start_column=1, end_row=7, end_column=2)
    ws.cell(7, 3, "No of Valid Votes Cast in favour of")
    ws.merge_cells(start_row=7, start_column=3, end_row=7, end_column=2 + nc)
    for j, t in enumerate(["Total of Valid Votes", "No. Of Rejected Votes", "Votes for NOTA", "Total",
                           "No. Of Tendered Votes"]):
        ws.cell(7, 3 + nc + j, t)
    for j, t in enumerate(["Serial No.", "Polling Station No"] + [n for n, _ in names]):
        ws.cell(8, 1 + j, t)
    for r in (7, 8):
        for c in range(1, ncols + 1):
            ws.cell(r, c).font, ws.cell(r, c).fill, ws.cell(r, c).alignment = bold, fill, center
    sums = [0] * (nc + 5)
    for n, (serial, ps_no, cand, tot) in enumerate(booths):
        if swap_nota:                              # the sheet prints NOTA before rejected; the extractor read them swapped
            tot = dict(tot, REJECTED=tot.get("NOTA"), NOTA=tot.get("REJECTED"))
        ps = int(ps_no) if str(ps_no).isdigit() else ps_no          # 41A stays text, like the existing files
        row = [serial, ps] + list(cand) + [tot.get(m) for m in TOTAL_METRICS]
        for j, v in enumerate(row):
            ws.cell(9 + n, 1 + j, v)
        for j, v in enumerate(row[2:]):
            sums[j] += v or 0
    last = 9 + len(booths)
    for k, label in enumerate(["Total EVM Votes", "Total Votes Polled"]):
        r = last + k
        ws.cell(r, 1, label).font = bold
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        for j, v in enumerate(sums):
            ws.cell(r, 3 + j, v).font = bold
    ws.freeze_panes = "A9"
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 10, 22
    for c in range(3, ncols + 1):
        ws.column_dimensions[L(c)].width = 16
    v = wb.create_sheet("Validation")
    for line in [
        f"Source: {source_name}",
        f"Booth rows extracted: {len(booths)}",
        f"Candidates ({nc}): " + ", ".join(f"{n} ({p})" if p else n for n, p in names),
        "Row-level check: sum(candidates)==Total Valid AND Valid+Rejected+NOTA==Total (OCR arithmetic repair applied)"
        + ("; NOTA and rejected columns were swapped in the OCR output and corrected using the official NOTA" if swap_nota
           else ""),
        "Official check: every candidate column and the NOTA total equal the official 2024 segment-level EVM totals"
        + ("" if verified else f" - NOT MET: {note}"),
        (f"Validation: PASS ({len(booths)} booth rows, all candidate columns + NOTA match the official EVM totals)"
         if verified else f"Validation: NOT VERIFIED ({len(booths)} booth rows; {note}) - check before use"),
    ]:
        v.append([line])
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main(argv=None):
    ap = argparse.ArgumentParser(description="Stage verified workbooks in the data/2024/EXCEL layout.")
    ap.add_argument("--ledger", default=str(QUEUE_DIR / "conversion_ledger.csv"))
    ap.add_argument("--check", default=str(QUEUE_DIR / "official_check.csv"))
    ap.add_argument("--out", default=str(DEFAULT_OUT))
    args = ap.parse_args(argv)
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    checks = {r["file"]: r for r in read_csv(args.check) if r["verdict"] == "FULL"}   # keyed by workbook name
    full = set(checks)
    official = load_official(OFFICIAL_CSV)
    staged, chosen = [], {}
    for r in sorted(read_csv(args.ledger), key=lambda r: (int(r["ac"]), int(r.get("repaired_rows") or 0))):
        if r["status"] == "DONE" and Path(r["xlsx"]).name in full and int(r["ac"]) not in chosen:
            chosen[int(r["ac"])] = r                           # fewest repaired rows first
    for ac, r in sorted(chosen.items()):
        labels, booths, _ = read_tidy(r["xlsx"])
        sums = [sum(b[2][i] or 0 for b in booths) for i in range(len(labels))]
        names = assign_names(labels, sums, official[ac]["candidates"])
        out = Path(args.out) / f"{ac}.xlsx"
        swap = checks[Path(r["xlsx"]).name].get("nota_swapped") == "yes"
        build_workbook(ac, r["file"], labels, booths, names, official[ac], out, swap_nota=swap)
        staged.append(ac)
    print(f"staged {len(staged)} workbooks in {args.out}: {staged}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
