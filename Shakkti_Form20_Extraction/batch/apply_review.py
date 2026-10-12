"""Apply the corrections downloaded from a review page, and check the result against the official totals.

    python batch/apply_review.py corrections_190.json            # writes data/2024/review/190_reviewed.xlsx and reports
    python batch/apply_review.py corrections_190.json --accept   # and files it into data/2024/converted/ when it matches

The AC's converted workbook is copied (the original is never changed); the corrected rows are written into the copy's Form20
sheet, a booth added by hand is inserted after the row it follows, and the copy is re-checked with batch/recheck_reviewed.py:
verified when every candidate column and NOTA equal the official 2024 totals.
"""
import argparse
import json
import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import openpyxl  # noqa: E402
from openpyxl.styles import PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter as L  # noqa: E402

from queue_common import REPO_ROOT  # noqa: E402
from recheck_reviewed import accept, recheck  # noqa: E402
from review_common import read_form20  # noqa: E402
from verify_official import load_official  # noqa: E402

DEFAULT_OUT = REPO_ROOT / "data" / "2024" / "review"
CORRECTED = PatternFill("solid", fgColor="C6EFCE")


def apply_corrections(workbook, corrections, out_path):
    """Write the corrections into a copy of `workbook` at out_path. Returns (rows changed, rows added)."""
    form = read_form20(workbook)
    nc, head = form["nc"], form["head"]
    where = {(b["page"], b["page_row"]): b["row"] for b in form["booths"]}
    shutil.copy2(workbook, out_path)
    wb = openpyxl.load_workbook(out_path)
    ws = wb["Form20"]
    names = {str(ws.cell(head, c).value): c for c in range(1, ws.max_column + 1) if ws.cell(head, c).value}
    col_note, col_page, col_prow = names["Review_Note"], names["Source_Page"], names["Source_Row"]

    def write(row, values, note):
        for j, v in enumerate(values):
            cell = ws.cell(row, 3 + j, None if v in ("", None) else int(v))
            cell.fill = CORRECTED
        ws.cell(row, col_note, note)

    changed = 0
    for e in corrections.get("edits", []):
        row = where.get((int(e["page"]), int(e["row"])))
        if row is None:
            raise SystemExit(f"page {e['page']} row {e['row']} is not in {workbook}")
        write(row, e["values"], "corrected by a person (review page)")
        changed += 1
    added = 0
    for a in sorted(corrections.get("added", []), key=lambda a: -where.get((int(a["page"]), int(a["after"])), 0)):
        after = where.get((int(a["page"]), int(a["after"])))
        if after is None:                                          # "0" = before the page's first row
            first = min((r for (p, _), r in where.items() if p == int(a["page"])), default=None)
            if first is None:
                raise SystemExit(f"page {a['page']} has no rows in {workbook}")
            after = first - 1
        ws.insert_rows(after + 1)
        write(after + 1, a["values"], "added by a person (review page): a booth the extractor missed")
        ws.cell(after + 1, col_page, int(a["page"]))
        ws.cell(after + 1, col_prow, f"added after {a['after']}")
        added += 1
    if added:                                                     # keep the printed 'Computed' SUM row pointing at the booth rows
        rows = [r for r in range(head + 2, ws.max_row + 1)
                if isinstance(ws.cell(r, col_page).value, (int, float)) or str(ws.cell(r, col_page).value or "").isdigit()]
        for r in range(1, ws.max_row + 1):
            if str(ws.cell(r, 1).value or "").startswith("Computed"):
                for c in range(3, 3 + nc + 5):
                    ws.cell(r, c, f"=SUM({L(c)}{min(rows)}:{L(c)}{max(rows)})")
    wb.save(out_path)
    return changed, added


def main(argv=None):
    ap = argparse.ArgumentParser(description="Apply review-page corrections to an AC's workbook and re-check it.")
    ap.add_argument("corrections")
    ap.add_argument("--accept", action="store_true", help="file the result into data/2024/converted/ when it matches")
    ap.add_argument("--out", default=str(DEFAULT_OUT))
    args = ap.parse_args(argv)
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    corrections = json.loads(Path(args.corrections).read_text(encoding="utf-8"))
    ac = int(corrections["ac"])
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    path = out / f"{ac}_reviewed.xlsx"
    changed, added = apply_corrections(corrections["workbook"], corrections, path)
    official = load_official()
    _, chk, form = recheck(path, official)
    print(f"AC {ac}: {changed} row(s) corrected, {added} added -> {path}")
    print(f"    {chk['verdict']}: {chk['matched']}/{chk['cols_official']} candidate columns exact, NOTA {chk['nota_sheet']} vs official "
          f"{chk['nota_official']}" + (f"\n    still off: {chk['detail']}" if chk["detail"] else ""))
    if chk["verdict"] == "FULL" and args.accept:
        print("    filed ->", accept(ac, path, form, chk, official))
    return 0


if __name__ == "__main__":
    sys.exit(main())
