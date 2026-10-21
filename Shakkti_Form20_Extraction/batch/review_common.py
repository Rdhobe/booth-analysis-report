"""Read a converted workbook's Form20 sheet back into rows (also after a person has corrected cells in it)."""
from pathlib import Path

import openpyxl

TOTAL_KEYS = ("valid", "rejected", "nota", "total", "tendered")


def read_form20(path):
    """Booth rows of the 'Form20' sheet of a workbook written by the extractor.

    Returns {"labels": candidate headers in sheet order, "nc": candidate count, "booths": [dict]}; each booth dict has
    row (sheet row), serial, ps_no, cands, valid, rejected, nota, total, tendered, status, note, page, page_row.
    The booth rows are the ones with a source page; the printed totals and the SUM row below them are not."""
    wb = openpyxl.load_workbook(path)
    try:
        ws = wb["Form20"]
        head = next(r for r in range(1, 30) if ws.cell(r, 1).value == "Serial No." and ws.cell(r, 2).value
                    and "Polling" in str(ws.cell(r, 2).value))
        top = head - 1                                          # the row above holds the group titles
        valid_col = next(c for c in range(3, ws.max_column + 1) if str(ws.cell(top, c).value or "").startswith("Total of"))
        nc = valid_col - 3
        names = {str(ws.cell(head, c).value): c for c in range(1, ws.max_column + 1) if ws.cell(head, c).value}
        col_status, col_note = names.get("Check_Status"), names.get("Review_Note")
        col_page, col_prow = names.get("Source_Page"), names.get("Source_Row")
        labels = [str(ws.cell(head, 3 + j).value or "") for j in range(nc)]
        booths = []
        for r in range(head + 2, ws.max_row + 1):
            page = ws.cell(r, col_page).value if col_page else None
            if not (isinstance(page, (int, float)) or str(page).isdigit()):
                continue
            vals = [ws.cell(r, valid_col + k).value for k in range(5)]
            booths.append({
                "row": r, "serial": ws.cell(r, 1).value, "ps_no": ws.cell(r, 2).value,
                "cands": [ws.cell(r, 3 + j).value for j in range(nc)],
                **dict(zip(TOTAL_KEYS, vals)),
                "status": ws.cell(r, col_status).value if col_status else None,
                "note": ws.cell(r, col_note).value if col_note else None,
                "page": int(page), "page_row": ws.cell(r, col_prow).value if col_prow else None,
            })
        title = [str(ws.cell(i, 1).value or "") for i in range(1, 7)]
        return {"labels": labels, "nc": nc, "booths": booths, "head": head, "title": title}
    finally:
        wb.close()


def column_sums(form):
    """Candidate column sums in sheet order [(label, sum)] and {"NOTA": n, "REJECTED": r} of a read_form20 result."""
    cand = [(form["labels"][j], sum((b["cands"][j] or 0) for b in form["booths"])) for j in range(form["nc"])]
    metrics = {"NOTA": sum((b["nota"] or 0) for b in form["booths"]),
               "REJECTED": sum((b["rejected"] or 0) for b in form["booths"])}
    return cand, metrics


def ledger_workbook(ac, ledger_rows, official, check_workbook_fn):
    """The converted workbook of an AC that is closest to the official totals (the PDF may have several)."""
    best = None
    for r in ledger_rows:
        if r["status"] != "DONE" or not r["xlsx"] or not Path(r["xlsx"]).exists():
            continue
        chk = check_workbook_fn(ac, r["xlsx"], official)
        key = (-(chk.get("matched") or 0), r["file"])
        if best is None or key < best[0]:
            best = (key, r, chk)
    return (best[1], best[2]) if best else (None, None)
