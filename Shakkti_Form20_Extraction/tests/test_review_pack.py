import csv
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "batch"))

from recheck_reviewed import accept, recheck  # noqa: E402
from review_common import column_sums, read_form20  # noqa: E402
from review_pack import classify, rank_pairs  # noqa: E402

OFFICIAL = {7: {"name": "Testpur", "nota": 6, "electors": 1000,
                "candidates": [("AAA", "P1", 90), ("BBB", "P2", 41), ("CCC", "P3", 9)]}}
BOOTHS = [(1, 1, [40, 20, 5], 3), (2, 2, [30, 11, 4], 1), (3, 3, [20, 10, 0], 2)]      # serial, ps, cands, NOTA


def make_workbook(path, drop_last=False):
    """A workbook in the extractor's Form20 layout: titles, two header rows, a Party row, booth rows, printed totals."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Form20"
    ws["A1"] = "FORM 20"
    ws["A8"], ws["C8"] = "Serial No. Of Polling Station", "No of Valid Votes Cast in favour of"
    ws.cell(8, 6, "Total of Valid Votes")
    for j, h in enumerate(["Serial No.", "Polling Station No", "AAA", "BBB", "CCC", "x1", "x2", "x3", "x4", "x5",
                           "Total Electors", "Check_Status", "Review_Note", "Source_Page", "Source_Row"], 1):
        ws.cell(9, j, h)
    ws.cell(10, 2, "Party")
    rows = BOOTHS[:-1] if drop_last else BOOTHS
    for i, (serial, ps, cands, nota) in enumerate(rows):
        r = 11 + i
        valid = sum(cands)
        for j, v in enumerate([serial, ps] + cands + [valid, 0, nota, valid + nota, 0]):
            ws.cell(r, 1 + j, v)
        ws.cell(r, 12, "OK" if i else "REVIEW")
        ws.cell(r, 14, 1)
        ws.cell(r, 15, i + 1)
    ws.cell(11 + len(rows), 1, "Total EVM votes")                    # printed totals row: no source page
    wb.create_sheet("Tidy")
    wb.save(path)


class ReviewPackTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_form20_rows_and_sums_are_read_back(self):
        make_workbook(self.dir / "7_review.xlsx")
        form = read_form20(self.dir / "7_review.xlsx")
        self.assertEqual((form["nc"], len(form["booths"])), (3, 3))
        self.assertEqual(form["labels"], ["AAA", "BBB", "CCC"])
        cand, metrics = column_sums(form)
        self.assertEqual([v for _, v in cand], [90, 41, 9])
        self.assertEqual(metrics["NOTA"], 6)
        self.assertEqual(form["booths"][0]["status"], "REVIEW")

    def test_matching_pack_is_verified_and_filed_with_the_manifest_updated(self):
        make_workbook(self.dir / "7_review.xlsx")
        conv = self.dir / "converted"
        conv.mkdir()
        with open(conv / "MANIFEST.csv", "w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=["ac", "status", "output_file", "how_produced", "check", "note_or_reason"])
            w.writeheader()
            w.writerow({"ac": 7, "status": "CONVERTED, NOT VERIFIED", "note_or_reason": "off"})
        ac, chk, form = recheck(self.dir / "7_review.xlsx", OFFICIAL)
        self.assertEqual((ac, chk["verdict"]), (7, "FULL"))
        out = accept(ac, self.dir / "7_review.xlsx", form, chk, OFFICIAL, conv)
        self.assertTrue(out.exists())
        row = next(csv.DictReader(open(conv / "MANIFEST.csv", encoding="utf-8-sig")))
        self.assertEqual((row["status"], row["output_file"]), ("VERIFIED", "converted/7.xlsx"))
        wb = openpyxl.load_workbook(out)
        self.assertIn("PASS", wb["Validation"]["A6"].value)

    def test_a_pack_with_a_missing_booth_is_not_verified(self):
        make_workbook(self.dir / "7_review.xlsx", drop_last=True)
        _, chk, _ = recheck(self.dir / "7_review.xlsx", OFFICIAL)
        self.assertNotEqual(chk["verdict"], "FULL")
        self.assertIn("-", chk["detail"])

    def test_columns_are_paired_with_official_candidates_by_rank(self):
        pairs = rank_pairs([("a", 41), ("b", 9), ("c", 90)], OFFICIAL[7]["candidates"])
        self.assertEqual(pairs, {1: ("CCC", 9), 0: ("BBB", 41), 2: ("AAA", 90)})

    def test_classes(self):
        base = dict(plan={}, unknown={}, est_missing=0.0, flagged=0, exact=3, n_cols=3, nota_diff=0)
        self.assertEqual(classify({"worst": 0.001}, "text layer", **dict(base, exact=1))[0], "SOURCE DIFFERS")
        self.assertEqual(classify({"worst": 0.02}, "ocr", **dict(base, plan={4: (7, 5)}, est_missing=11.0, exact=0))[0], "ROWS MISSING")
        self.assertEqual(classify({"worst": 0.02}, "ocr", **dict(base, est_missing=-3.0, exact=0))[0], "EXTRA ROWS")
        self.assertEqual(classify({"worst": 0.0}, "text layer", **dict(base, nota_diff=700))[0], "NOTA COLUMN")
        self.assertEqual(classify({"worst": 0.02}, "ocr", **dict(base, exact=0))[0], "CHECK ROWS")
        self.assertEqual(classify({"worst": 0.4}, "ocr", **dict(base, exact=0))[0], "NEEDS CODE FIX")


if __name__ == "__main__":
    unittest.main()
