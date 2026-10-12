import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "batch"))

from apply_review import apply_corrections  # noqa: E402
from recheck_reviewed import recheck  # noqa: E402
from review_common import read_form20  # noqa: E402
from review_page import build_data, read_boxes, render_html, select_rows  # noqa: E402
from tests.test_review_pack import BOOTHS, OFFICIAL, make_workbook  # noqa: E402


def with_boxes(path, rows):
    """Add a Row_Boxes sheet: (page, source_row, serial, page_index, x0, y0, x1, y1, plain, changed)."""
    wb = openpyxl.load_workbook(path)
    ws = wb.create_sheet("Row_Boxes")
    ws.append(["Page", "Source_Row", "Serial", "Page_Index", "X0", "Y0", "X1", "Y1", "Plain_Reading", "Changed"])
    for r in rows:
        ws.append(list(r))
    wb.save(path)


class ReviewPageTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.path = self.dir / "7_review.xlsx"
        make_workbook(self.path)                                 # row 1 is REVIEW, rows 2 and 3 are OK
        with_boxes(self.path, [(1, 1, 1, 0, 10, 100, 900, 140, "40 | 20 | 5 | 65 | 0 | 3 | 68 | 0", "0"),
                               (1, 2, 2, 0, 10, 140, 900, 180, "30 | 11 | 4 | 45 | 0 | 1 | 46 | 0", ""),
                               (1, 3, 3, 0, 10, 180, 900, 220, "20 | 10 | 0 | 30 | 0 | 2 | 32 | 0", "")])

    def tearDown(self):
        self.tmp.cleanup()

    def test_only_flagged_rows_with_a_box_are_shown(self):
        form, boxes = read_form20(self.path), read_boxes(self.path)
        self.assertEqual([b["serial"] for b in select_rows(form["booths"], boxes, 10)], [1])

    def test_an_auto_corrected_row_counts_only_if_a_read_cell_was_changed(self):
        form, boxes = read_form20(self.path), read_boxes(self.path)
        form["booths"][1]["status"] = "AUTO_CORRECTED"
        boxes[(1, 2)]["plain"] = ["30", "11", "4", "45", "0", "1", "", "0"]
        boxes[(1, 2)]["changed"] = [6]                                # the changed cell was blank: a fill, not a doubt
        self.assertEqual([b["serial"] for b in select_rows(form["booths"], boxes, 10)], [1])
        boxes[(1, 2)]["changed"] = [0]                                # a cell the OCR had read was changed
        self.assertEqual([b["serial"] for b in select_rows(form["booths"], boxes, 10)], [1, 2])

    def test_the_limit_keeps_the_most_doubtful_rows(self):
        form, boxes = read_form20(self.path), read_boxes(self.path)
        for b in form["booths"][1:]:
            b["status"] = "UNUSUAL"
        self.assertEqual([b["serial"] for b in select_rows(form["booths"], boxes, 1)], [1])       # REVIEW outranks UNUSUAL

    def test_the_page_holds_the_base_totals_of_the_other_rows_and_the_official_ones(self):
        form, boxes = read_form20(self.path), read_boxes(self.path)
        data = build_data(7, OFFICIAL[7], form, boxes, self.path, "7Eng.pdf", 10, lambda info: "data:image/jpeg;base64,AAAA")
        self.assertEqual(sorted(c["official"] for c in data["candidates"]), [9, 41, 90])
        shown = data["rows"][0]
        self.assertEqual((shown["serial"], shown["status"], shown["img"][:10]), (1, "REVIEW", "data:image"))
        for j, c in enumerate(data["candidates"]):                    # the rows that are not shown plus the shown one = the sheet's sum
            self.assertEqual(c["base"] + shown["values"][j], sum(b[2][j] for b in BOOTHS))
        self.assertEqual(data["nota"]["base"] + shown["values"][len(data["candidates"]) + 2], sum(b[3] for b in BOOTHS))

    def test_the_html_carries_the_data_and_cannot_be_broken_by_text_in_it(self):
        form, boxes = read_form20(self.path), read_boxes(self.path)
        form["booths"][0]["note"] = "</script><b>x"
        data = build_data(7, OFFICIAL[7], form, boxes, self.path, "7Eng.pdf", 10, lambda info: "")
        html = render_html(data)
        self.assertIn("const DATA = {", html)
        self.assertNotIn("</script><b>", html)
        self.assertEqual(html.count("</script>"), 1)
        self.assertIn("Review AC 7", html)

    def test_corrections_are_applied_to_a_copy_and_re_checked(self):
        before = read_form20(self.path)
        right = before["booths"][0]["cands"][:]
        wb = openpyxl.load_workbook(self.path)                        # break the sheet first: a misread cell in row 1
        wb["Form20"].cell(before["booths"][0]["row"], 3, 41)
        wb.save(self.path)
        _, chk, _ = recheck(self.path, OFFICIAL)
        self.assertNotEqual(chk["verdict"], "FULL")
        out = self.dir / "7_reviewed.xlsx"
        corrections = {"ac": 7, "workbook": str(self.path),
                       "edits": [{"page": 1, "row": 1, "values": right + [sum(right), 0, 3, sum(right) + 3, 0]}], "added": []}
        self.assertEqual(apply_corrections(self.path, corrections, out), (1, 0))
        _, chk, _ = recheck(out, OFFICIAL)
        self.assertEqual(chk["verdict"], "FULL")
        _, chk_original, _ = recheck(self.path, OFFICIAL)
        self.assertNotEqual(chk_original["verdict"], "FULL")           # the original is untouched

    def test_a_booth_added_by_hand_is_inserted_after_its_row(self):
        form = read_form20(self.path)
        wb = openpyxl.load_workbook(self.path)                        # drop the last booth ...
        wb["Form20"].delete_rows(form["booths"][2]["row"])
        wb.save(self.path)
        _, chk, _ = recheck(self.path, OFFICIAL)
        self.assertNotEqual(chk["verdict"], "FULL")
        last = BOOTHS[2]                                               # ... and add it back through the page
        values = last[2] + [sum(last[2]), 0, last[3], sum(last[2]) + last[3], 0]
        corrections = {"ac": 7, "workbook": str(self.path), "edits": [],
                       "added": [{"page": 1, "after": 2, "values": values}]}
        out = self.dir / "7_reviewed.xlsx"
        self.assertEqual(apply_corrections(self.path, corrections, out), (0, 1))
        _, chk, form2 = recheck(out, OFFICIAL)
        self.assertEqual(chk["verdict"], "FULL")
        self.assertEqual(len(form2["booths"]), 3)

    def test_a_row_marked_not_a_booth_is_deleted_from_the_copy(self):
        before = read_form20(self.path)
        out = self.dir / "7_reviewed.xlsx"
        corrections = {"ac": 7, "workbook": str(self.path), "edits": [], "added": [], "removed": [{"page": 1, "row": 3}]}
        self.assertEqual(apply_corrections(self.path, corrections, out), (0, 0))
        after = read_form20(out)
        self.assertEqual(len(after["booths"]), len(before["booths"]) - 1)
        self.assertEqual([b["serial"] for b in after["booths"]], [b["serial"] for b in before["booths"][:2]])
        _, chk, _ = recheck(out, OFFICIAL)
        self.assertNotEqual(chk["verdict"], "FULL")                    # its votes are gone from the totals
        self.assertEqual(len(read_form20(self.path)["booths"]), len(before["booths"]))      # the original is untouched

    def test_removing_a_row_that_is_not_there_stops_with_a_message(self):
        corrections = {"ac": 7, "workbook": str(self.path), "edits": [], "added": [], "removed": [{"page": 9, "row": 9}]}
        with self.assertRaises(SystemExit):
            apply_corrections(self.path, corrections, self.dir / "7_reviewed.xlsx")

    def test_the_page_offers_to_remove_a_row_and_leaves_it_out_of_the_totals(self):
        form, boxes = read_form20(self.path), read_boxes(self.path)
        data = build_data(7, OFFICIAL[7], form, boxes, self.path, "7Eng.pdf", 10, lambda info: "")
        html = render_html(data)
        for needle in ("Not a booth", "state.removed", "removed }", ".filter(r => !state.removed[r.key])"):
            self.assertIn(needle, html)

    def test_booths_added_on_two_pages_land_where_they_belong_whatever_the_order_given(self):
        wb = openpyxl.load_workbook(self.path)                         # row 3 moves to page 2 (its first row)
        ws = wb["Form20"]
        ws.cell(13, 14, 2)
        ws.cell(13, 15, 1)
        wb.save(self.path)
        with_boxes(self.path, [(1, 1, 1, 0, 10, 100, 900, 140, "", ""), (1, 2, 2, 0, 10, 140, 900, 180, "", ""),
                               (2, 1, 3, 1, 10, 100, 900, 140, "", "")])
        mark = lambda k: [k, 0, 0, k, 0, 0, k, 0]                      # three candidates, then valid, rejected, NOTA, total, tendered
        corrections = {"ac": 7, "workbook": str(self.path), "edits": [], "removed": [],
                       "added": [{"page": 2, "after": 0, "values": mark(8)},       # before the first booth of page 2
                                 {"page": 1, "after": 2, "values": mark(7)}]}      # after the last booth of page 1
        out = self.dir / "7_reviewed.xlsx"
        self.assertEqual(apply_corrections(self.path, corrections, out), (0, 2))
        got = [(b["page"], b["cands"][0]) for b in read_form20(out)["booths"]]
        self.assertEqual(got, [(1, 40), (1, 30), (1, 7), (2, 8), (2, 20)])      # page 1, page 1, added to 1, added to 2, page 2's own


if __name__ == "__main__":
    unittest.main()
