import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "batch"))

import openpyxl  # noqa: E402

from stage_verified import assign_names, build_workbook  # noqa: E402
from verify_official import compare, sheet_sums  # noqa: E402


class AssignNamesTests(unittest.TestCase):
    OFFICIAL = [("ASHA DEVI", "BJP", 500), ("RAM KUMAR", "SP", 300), ("SITA", "IND", 0), ("GITA", "IND", 0)]

    def test_columns_are_matched_to_official_names_by_total(self):
        labels = ["RAM KUMA", "ASHA DEV", "Candidate 3", "Candidate 4"]
        got = assign_names(labels, [300, 500, 0, 0], self.OFFICIAL)
        self.assertEqual([n for n, _ in got][:2], ["RAM KUMAR", "ASHA DEVI"])
        self.assertEqual({n for n, _ in got[2:]}, {"SITA", "GITA"})            # equal totals: both used once

    def test_equal_totals_are_told_apart_by_name_similarity(self):
        got = assign_names(["GITA D", "SITA"], [0, 0], self.OFFICIAL[2:])
        self.assertEqual([n for n, _ in got], ["GITA", "SITA"])

    def test_unmatched_column_keeps_its_ocr_label(self):
        got = assign_names(["WEIRD"], [999], self.OFFICIAL)
        self.assertEqual(got, [("WEIRD", "")])


class SheetSumsTests(unittest.TestCase):
    def test_columns_with_the_same_name_are_not_merged(self):
        import openpyxl
        from verify_official import sheet_sums
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Tidy"
        ws.append(["Source_File", "Segment_No", "Segment_Name", "Electors", "PS_Serial", "PS_No", "Metric", "Name",
                   "Name_Native", "Party", "Votes", "Check_Status"])
        for serial, (a, b, nota) in enumerate([(5, 3, 1), (7, 2, 0)], 1):          # two columns, both named "X"
            for v in (a, b):
                ws.append(["f", None, None, None, serial, str(serial), "CANDIDATE", "X", "X", None, v, "OK"])
            ws.append(["f", None, None, None, serial, str(serial), "NOTA", "Votes for NOTA", "", None, nota, "OK"])
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "w.xlsx"
            wb.save(path)
            cand, metrics, missing = sheet_sums(path)
        self.assertEqual([t for _, t in cand], [12, 5])                           # not one merged column of 17
        self.assertEqual(metrics["NOTA"], 1)
        self.assertEqual(missing, 0)


class BuildWorkbookTests(unittest.TestCase):
    def test_workbook_follows_the_existing_excel_layout(self):
        labels = ["A", "B"]
        booths = [(1, "1", [5, 3], {"TOTAL_VALID": 8, "REJECTED": 0, "NOTA": 1, "TOTAL": 9, "TENDERED": 0}),
                  (2, "2", [7, 2], {"TOTAL_VALID": 9, "REJECTED": 1, "NOTA": 0, "TOTAL": 10, "TENDERED": 0})]
        names = [("ASHA DEVI", "BJP"), ("RAM KUMAR", "SP")]
        official = {"name": "Rath", "electors": 409486, "nota": 1, "candidates": [("ASHA DEVI", "BJP", 12), ("RAM KUMAR", "SP", 5)]}
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "229.xlsx"
            build_workbook(229, "229Eng.pdf", labels, booths, names, official, out)
            wb = openpyxl.load_workbook(out)
        ws = wb["Form20"]
        self.assertEqual(ws["A1"].value, "FORM 20")
        self.assertIn("....409486", ws["A4"].value)
        self.assertIn("229-Rath Assembly Election", ws["A5"].value)
        self.assertEqual([ws.cell(8, c).value for c in range(1, 5)], ["Serial No.", "Polling Station No", "ASHA DEVI", "RAM KUMAR"])
        self.assertEqual([ws.cell(9, c).value for c in range(1, 10)], [1, 1, 5, 3, 8, 0, 1, 9, 0])   # numbers, not text
        self.assertEqual(ws.cell(11, 1).value, "Total EVM Votes")
        self.assertEqual([ws.cell(11, c).value for c in range(3, 5)], [12, 5])              # column sums
        self.assertEqual(ws.cell(11, 7).value, 1)                                            # NOTA total
        self.assertEqual(ws.freeze_panes, "A9")
        self.assertIn("PASS", wb["Validation"]["A6"].value)

    def test_swapped_nota_and_rejected_are_corrected_in_the_staged_file(self):
        booths = [(1, "1", [5, 3], {"TOTAL_VALID": 8, "REJECTED": 4, "NOTA": 0, "TOTAL": 12, "TENDERED": 0})]   # read swapped
        official = {"name": "Rath", "electors": 1, "nota": 4, "candidates": [("A", "", 5), ("B", "", 3)]}
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "1.xlsx"
            build_workbook(1, "1Hin.pdf", ["x", "y"], booths, [("A", ""), ("B", "")], official, out, swap_nota=True)
            wb = openpyxl.load_workbook(out)
            ws = wb["Form20"]
            row = [ws.cell(9, c).value for c in range(1, 10)]
        self.assertEqual(row, [1, 1, 5, 3, 8, 0, 4, 12, 0])        # rejected 0, NOTA 4
        self.assertIn("swapped", wb["Validation"]["A4"].value)

    def test_written_workbook_still_passes_the_official_comparison(self):
        cand = [("ASHA DEVI", 12), ("RAM KUMAR", 5)]
        official = {"nota": 1, "candidates": [("ASHA DEVI", "BJP", 12), ("RAM KUMAR", "SP", 5)]}
        self.assertEqual(compare(cand, {"NOTA": 1}, official)["verdict"], "FULL")


if __name__ == "__main__":
    unittest.main()
